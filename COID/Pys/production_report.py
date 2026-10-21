from __future__ import annotations

import argparse
import logging
import math
import re
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
COID_DIR = HERE.parent
EXPORT_DIR = COID_DIR / "Export"
SHEET_PREFIX = "PlanVSAchvd_"
LOG = logging.getLogger("COID_REPORT")
FAMILY_ORDER = ["HALB", "Standard Cubes", "Double Cubes", "Mixpy", "Signature", "Tablets", "Pouches"]


def product_family(material_type: str, description: str) -> str:
    if material_type == "HALB":
        return "HALB"
    name = description.upper()
    if "DOUBLE CUBE" in name:
        return "Double Cubes"
    if "MIXPY" in name:
        return "Mixpy"
    if "SIGNATURE" in name:
        return "Signature"
    if "TABLET" in name:
        return "Tablets"
    if "CUBE" in name:
        return "Standard Cubes"
    return "Pouches"


def pack_weight_grams(description: str) -> float | None:
    name = description.upper().replace(",", "")
    patterns = (
        r"(\d+)\s*[X×]\s*(\d+)\s*[X×]\s*(\d+(?:\.\d+)?)\s*G",
        r"(\d+)\s*\(\s*(\d+)\s*[X×]\s*(\d+(?:\.\d+)?)\s*G\s*\)",
        r"(\d+)\s*[X×]\s*(\d+(?:\.\d+)?)\s*G",
    )
    for pattern in patterns:
        match = re.search(pattern, name)
        if match:
            return math.prod(float(value) for value in match.groups())
    return None


def quantity_tons(row: pd.Series) -> tuple[float | None, str]:
    if str(row["Order Type"]).strip().upper() != "PI01":
        return None, "Excluded order type (chart includes PI01 production orders only)"
    unit = str(row["Unit of measure"]).strip().upper()
    if row["Material Type"] == "HALB" and unit == "KG":
        return 0.001, ""
    if row["Material Type"] == "FERT" and unit == "CS":
        weight = pack_weight_grams(str(row["Material Description"]))
        if weight is not None:
            return weight / 1_000_000, ""
        return None, "No pack weight found in material description"
    return None, f"Unsupported unit/material type: {row['Material Type']} {unit}"


def build_report(
    input_path: Path,
    output_path: Path | None = None,
    period_start: date | None = None,
    period_end: date | None = None,
) -> Path:
    LOG.info("STEP report read input: %s", input_path)
    frame = pd.read_csv(input_path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
    required = [
        "Material", "Material Description", "Basic Start Date", "Delivered Quantity",
        "Planned Quantity", "Order Type", "Unit of measure",
    ]
    missing = [column for column in required if column not in frame.columns]
    if missing:
        raise RuntimeError(f"Cannot build report; missing columns: {missing}")
    LOG.info("Report input validated: rows=%d", len(frame))
    frame["Basic Start Date"] = pd.to_datetime(
        frame["Basic Start Date"], format="mixed", errors="coerce"
    ).dt.strftime("%Y-%m-%d")
    if period_start:
        frame = frame[frame["Basic Start Date"] >= period_start.isoformat()].copy()
    if period_end:
        frame = frame[frame["Basic Start Date"] <= period_end.isoformat()].copy()
    frame["Delivered Quantity"] = pd.to_numeric(frame["Delivered Quantity"], errors="coerce").fillna(0)
    frame["Planned Quantity"] = pd.to_numeric(frame["Planned Quantity"], errors="coerce").fillna(0)
    frame["Material"] = frame["Material"].astype(str).str.strip()
    frame["Material Type"] = frame["Material"].str[0].map({"4": "HALB", "1": "FERT"}).fillna("Other")
    frame["Product Family"] = [
        product_family(material_type, description)
        for material_type, description in zip(frame["Material Type"], frame["Material Description"])
    ]
    conversions = frame.apply(quantity_tons, axis=1)
    frame["Tons per Unit"] = [value for value, _ in conversions]
    frame["Exception Reason"] = [reason for _, reason in conversions]
    eligible = frame[frame["Exception Reason"].eq("")].copy()
    if eligible.empty:
        raise RuntimeError("No in-period PI01 rows have supported units for the tons report")
    eligible["Planned Tons"] = eligible["Planned Quantity"] * eligible["Tons per Unit"]
    eligible["Achieved Tons"] = eligible["Delivered Quantity"] * eligible["Tons per Unit"]
    sort_rank = {family: index for index, family in enumerate(FAMILY_ORDER)}
    eligible["_Family Sort"] = eligible["Product Family"].map(sort_rank).fillna(len(sort_rank))
    eligible = eligible.sort_values(
        ["_Family Sort", "Material Description", "Material", "Basic Start Date"], kind="stable"
    )
    pivot = eligible.groupby(
        ["Product Family", "Material Type", "Material", "Material Description", "Unit of measure"],
        as_index=False,
        sort=False,
    )[["Planned Tons", "Achieved Tons"]].sum()
    pivot["_Family Sort"] = pivot["Product Family"].map(sort_rank).fillna(len(sort_rank))
    pivot = pivot.sort_values(
        ["_Family Sort", "Material Description", "Material"], kind="stable"
    ).drop(columns="_Family Sort")
    chart_summary = eligible.groupby(
        ["Product Family", "Material Description"], as_index=False, sort=False
    )[["Planned Tons", "Achieved Tons"]].sum()
    chart_summary["_Family Sort"] = chart_summary["Product Family"].map(sort_rank).fillna(len(sort_rank))
    chart_summary = chart_summary.sort_values(
        ["_Family Sort", "Material Description"], kind="stable"
    )
    active_families = [
        family
        for family in FAMILY_ORDER
        if chart_summary.loc[chart_summary["Product Family"].eq(family), "Planned Tons"].sum() > 0
    ]
    chart_data_parts = []
    for family in active_families:
        section = chart_summary[chart_summary["Product Family"].eq(family)]
        chart_data_parts.append(pd.DataFrame({
            f"{family} Product": section["Material Description"].reset_index(drop=True),
            f"{family} Planned (t)": section["Planned Tons"].reset_index(drop=True),
            f"{family} Achieved (t)": section["Achieved Tons"].reset_index(drop=True),
        }))
    chart_data = pd.concat(chart_data_parts, axis=1)
    exceptions = frame[frame["Exception Reason"].ne("")].copy()
    scope_start = period_start.isoformat() if period_start else "All dates"
    scope_end = period_end.isoformat() if period_end else "All dates"
    if output_path is None:
        output_path = EXPORT_DIR / f"Production Plan_{date.today():%Y-%m-%d}.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pivot.to_excel(writer, sheet_name="WTD_Pivot", index=False)
        frame.to_excel(writer, sheet_name="Source_Data", index=False)
        chart_data.to_excel(writer, sheet_name="Chart_Data", index=False)
        exceptions.to_excel(writer, sheet_name="Exceptions", index=False)
    LOG.info("STEP report workbook written: %s", output_path)
    format_workbook(output_path, scope_start, scope_end, len(exceptions), active_families)
    LOG.info("STEP report workbook formatted and chart added: %s", output_path)
    return output_path


def format_workbook(
    path: Path,
    scope_start: str,
    scope_end: str,
    exception_count: int,
    active_families: list[str],
) -> None:
    LOG.info("STEP format workbook: %s", path)
    book = load_workbook(path)
    pivot = book["WTD_Pivot"]
    pivot.freeze_panes = "A2"
    pivot.auto_filter.ref = f"A1:{get_column_letter(pivot.max_column)}{pivot.max_row}"
    for cell in pivot[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column in pivot.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
        pivot.column_dimensions[get_column_letter(column[0].column)].width = width
    source = book["Source_Data"]
    source.freeze_panes = "A2"
    source.auto_filter.ref = source.dimensions
    for cell in source[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    exceptions = book["Exceptions"]
    exceptions.freeze_panes = "A2"
    exceptions.auto_filter.ref = exceptions.dimensions
    for cell in exceptions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="9C3B2E")
    dashboard = book.create_sheet("Dashboard", 0)
    dashboard.merge_cells("A1:P1")
    dashboard["A1"] = "PRODUCTION PLAN vs ACHIEVED (WTD) (Tons)"
    dashboard["A1"].font = Font(name="Arial", size=20, bold=True, color="000000")
    dashboard["A1"].fill = PatternFill("solid", fgColor="A8C8E8")
    dashboard["A1"].alignment = Alignment(horizontal="center")
    dashboard.row_dimensions[1].height = 34
    dashboard.merge_cells("A2:P2")
    dashboard["A2"] = (
        f"Period: {scope_start} to {scope_end} | PI01 production orders only | "
        f"HALB kilograms divided by 1,000; FERT case weight from SKU grams / 1,000,000 | Exceptions: {exception_count}"
    )
    dashboard["A2"].font = Font(size=10, italic=True, color="404040")
    chart_data = book["Chart_Data"]
    if len(active_families) <= 6:
        positions = ["A4", "AG4", "A22", "O22", "AI22", "AW22"]
        chart_widths = [32, 24, 14, 20, 14, 18]
    else:
        positions = ["A4", "AG4", "A22", "X22", "AU22", "A40", "AG40"]
        chart_widths = [32, 32, 20, 20, 20, 32, 32]
    colors = ["176B87", "39802D"]
    for index, family in enumerate(active_families):
        product_column = index * 3 + 1
        first_data_row = 2
        last_data_row = first_data_row
        while (
            last_data_row <= chart_data.max_row
            and chart_data.cell(last_data_row, product_column).value is not None
        ):
            last_data_row += 1
        last_data_row -= 1
        if last_data_row < first_data_row:
            continue
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        title = "HALB" if family == "HALB" else f"FERT ({family})"
        chart.title = f"{title} - Plan Vs Achieved"
        chart.y_axis.title = "Tons"
        chart.x_axis.title = "Product"
        chart.legend.position = "b"
        chart.style = 10
        chart.add_data(
            Reference(chart_data, min_col=product_column + 1, max_col=product_column + 2,
                      min_row=1, max_row=last_data_row),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(chart_data, min_col=product_column, min_row=first_data_row,
                      max_row=last_data_row)
        )
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.height = 9
        chart.width = chart_widths[index]
        for series, color in zip(chart.series, colors):
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        dashboard.add_chart(chart, positions[index])
    dashboard.sheet_view.showGridLines = False
    dashboard.freeze_panes = "A3"
    book.save(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the COID plan-versus-achieved workbook.")
    parser.add_argument("--input", type=Path, default=COID_DIR / "CleanFile" / "COID_4981_clean.csv")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    print(build_report(args.input, args.output))
