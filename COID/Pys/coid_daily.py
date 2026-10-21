from __future__ import annotations

import argparse
import logging
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
COID_DIR = HERE.parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from coid_engine import reporting_period, run  # noqa: E402
from production_report import build_report  # noqa: E402


def report_is_complete(path: Path, period_start: date, period_end: date) -> bool:
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            required = {"Dashboard", "WTD_Pivot", "Breakdown"}
            if not required.issubset(workbook.sheetnames):
                return False
            dashboard = workbook["Dashboard"]
            period_text = f"Period: {period_start.isoformat()} to {period_end.isoformat()}"
            return (
                period_text in str(dashboard["A2"].value)
                and bool(dashboard._charts)
                and len(workbook["WTD_Pivot"]._pivots) == 1
                and workbook["Breakdown"].max_row > 1
            )
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError):
        return False


def main(skip_if_valid_report: bool = False) -> int:
    log = logging.getLogger("COID_DAILY")
    log.info("STEP daily workflow start")
    try:
        run_date = date.today()
        period_start, period_end, report_tag = reporting_period(run_date)
        report_path = COID_DIR / "Export" / f"Production Plan_{report_tag}.xlsx"
        if skip_if_valid_report and run_date.weekday() >= 5:
            print("COID weekday automation; skipping weekend wake event.")
            return 0
        if skip_if_valid_report and report_is_complete(report_path, period_start, period_end):
            print(f"COID report already complete for {period_start}..{period_end}; skipping SAP run.")
            return 0
        log.info("STEP extract and normalize COID")
        clean_file = run(dry_run=False, today=run_date)
        if clean_file is None:
            raise RuntimeError("COID did not produce a clean file")
        log.info("STEP build Excel report from %s", clean_file)
        report = build_report(
            clean_file,
            output_path=report_path,
            period_start=period_start,
            period_end=period_end,
        )
        log.info("STEP daily workflow complete: %s", report)
        print(f"COID daily workbook: {report}")
        return 0
    except Exception:
        log.exception("Daily COID workflow failed")
        raise


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Run the COID SAP extraction and WTD report workflow.")
    parser.add_argument(
        "--skip-if-valid-report",
        action="store_true",
        help="Skip SAP extraction when a complete report already exists for this reporting period.",
    )
    args = parser.parse_args()
    raise SystemExit(main(skip_if_valid_report=args.skip_if_valid_report))
